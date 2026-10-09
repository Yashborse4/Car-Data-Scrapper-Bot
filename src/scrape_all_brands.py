"""
Spinny Car Data Scraper - ALL BRANDS Complete Scrape
=====================================================

Scrapes ALL usable brands from Spinny API (31+ brands) for years 2000-2025.

Usage:
    python scrape_all_brands.py
"""

import os
import sys
import json
import time
import random
import logging
from datetime import datetime
from typing import Dict, List, Optional, Any

import requests
import pandas as pd
from openpyxl.styles import Font, Alignment, PatternFill
from tenacity import retry, stop_after_attempt, wait_exponential, retry_if_exception_type
from tqdm import tqdm


# =============================================================================
# CONFIGURATION
# =============================================================================

BASE_URL = "https://api.spinny.com"
MAKE_LIST_ENDPOINT = "/v3/api/catalogue/make-list"
MODEL_LIST_ENDPOINT = "/v3/api/catalogue/model-list"
VARIANT_LIST_ENDPOINT = "/v3/api/catalogue/variant-list"

HEADERS = {
    "platform": "mweb_android",
    "user-agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36",
    "accept": "application/json",
    "origin": "https://www.spinny.com",
    "referer": "https://www.spinny.com/",
}

# Rate limiting
REQUEST_DELAY = 1.2  # seconds between requests
JITTER = 0.3  # random jitter

# Scrape range
YEAR_START = 2000
YEAR_END = 2025

# Output
OUTPUT_DIR = "output"
CHECKPOINT_FILE = "all_brands_checkpoint.json"

# Setup logging
os.makedirs(OUTPUT_DIR, exist_ok=True)
logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s - %(levelname)s - %(message)s",
    handlers=[
        logging.StreamHandler(sys.stdout),
        logging.FileHandler(f"{OUTPUT_DIR}/all_brands_scraper.log", encoding='utf-8')
    ]
)
logger = logging.getLogger(__name__)


# =============================================================================
# API CLIENT
# =============================================================================

class SpinnyClient:
    def __init__(self):
        self.session = requests.Session()
        self.session.headers.update(HEADERS)
        self.last_request = 0
        self.request_count = 0
    
    def _rate_limit(self):
        elapsed = time.time() - self.last_request
        delay = REQUEST_DELAY + random.uniform(-JITTER, JITTER)
        if elapsed < delay:
            time.sleep(delay - elapsed)
        self.last_request = time.time()
        self.request_count += 1
    
    @retry(
        stop=stop_after_attempt(5),
        wait=wait_exponential(multiplier=2, max=60),
        retry=retry_if_exception_type(requests.RequestException),
    )
    def get(self, endpoint: str, params: Optional[Dict] = None) -> Dict:
        self._rate_limit()
        url = f"{BASE_URL}{endpoint}"
        resp = self.session.get(url, params=params, timeout=30)
        
        if resp.status_code == 429:
            logger.warning("Rate limit hit, backing off...")
            time.sleep(30)
            raise requests.RequestException("Rate limit")
        
        resp.raise_for_status()
        return resp.json()
    
    def get_all_brands(self) -> List[Dict]:
        """Get ALL brands (both usable and non-usable for reference)."""
        resp = self.get(MAKE_LIST_ENDPOINT)
        return resp.get("results", [])
    
    def get_models(self, make_id: int, year: int) -> List[Dict]:
        resp = self.get(MODEL_LIST_ENDPOINT, {"make_id": make_id, "make_year": year})
        return resp.get("results", [])
    
    def get_variants(self, model_id: int, year: int) -> List[Dict]:
        resp = self.get(VARIANT_LIST_ENDPOINT, {"model_id": model_id, "make_year": year})
        return resp.get("results", [])
    
    def close(self):
        self.session.close()


# =============================================================================
# DATA EXTRACTION
# =============================================================================

def extract_attr(variant: Dict, name: str) -> Any:
    """Extract attribute by display_name (not index)."""
    for attr in variant.get("attributes", []):
        if name.lower() in attr.get("display_name", "").lower():
            return attr.get("value")
    return None


def normalize(variant: Dict, brand: str, model: str, body_type: str, segment: str, year: int) -> Dict:
    """Normalize variant to flat dict."""
    return {
        "variant_id": variant.get("id"),
        "brand_name": brand,
        "model_name": model,
        "model_id": variant.get("model"),
        "variant_name": variant.get("display_name", ""),
        "variant_full_name": variant.get("full_name", ""),
        "fuel_type": variant.get("fuel_type", ""),
        "transmission": variant.get("transmission_type", ""),
        "body_type": body_type,
        "segment": segment,
        "engine_cc": extract_attr(variant, "Displacement"),
        "mileage_arai": extract_attr(variant, "Mileage"),
        "power_hp": extract_attr(variant, "Power"),
        "torque_nm": extract_attr(variant, "Torque"),
        "launch_year": variant.get("launch_year"),
        "discontinued_year": variant.get("discontinued_year"),
        "make_year": year,
        "is_usable": variant.get("is_usable", False),
        "is_active": variant.get("is_active", False),
    }


# =============================================================================
# CHECKPOINT MANAGER
# =============================================================================

def load_checkpoint() -> Optional[Dict]:
    path = os.path.join(OUTPUT_DIR, CHECKPOINT_FILE)
    if os.path.exists(path):
        with open(path, 'r', encoding='utf-8') as f:
            return json.load(f)
    return None


def save_checkpoint(rows: List, brand_idx: int, year_idx: int, errors: List):
    path = os.path.join(OUTPUT_DIR, CHECKPOINT_FILE)
    with open(path, 'w', encoding='utf-8') as f:
        json.dump({
            "rows": rows,
            "brand_idx": brand_idx,
            "year_idx": year_idx,
            "errors": errors,
            "timestamp": datetime.now().isoformat()
        }, f, ensure_ascii=False)


def delete_checkpoint():
    path = os.path.join(OUTPUT_DIR, CHECKPOINT_FILE)
    if os.path.exists(path):
        os.remove(path)


# =============================================================================
# EXCEL EXPORT
# =============================================================================

def export_excel(rows: List[Dict], errors: List[Dict]) -> str:
    if not rows:
        logger.warning("No data to export")
        return ""
    
    df = pd.DataFrame(rows)
    
    # Column order
    cols = [
        "variant_id", "brand_name", "model_name", "variant_name", "variant_full_name",
        "fuel_type", "transmission", "body_type", "segment", "engine_cc", "mileage_arai",
        "power_hp", "torque_nm", "launch_year", "discontinued_year", "make_year",
        "is_usable", "is_active", "model_id"
    ]
    df = df[[c for c in cols if c in df.columns]]
    
    # Remove duplicates by variant_id
    before = len(df)
    df = df.drop_duplicates(subset=["variant_id"], keep="first")
    logger.info(f"Removed {before - len(df)} duplicates")
    
    # Sort
    df = df.sort_values(["brand_name", "model_name", "make_year", "variant_name"])
    
    # Export
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    path = os.path.join(OUTPUT_DIR, f"spinny_all_brands_{timestamp}.xlsx")
    
    with pd.ExcelWriter(path, engine='openpyxl') as writer:
        df.to_excel(writer, sheet_name='All Variants', index=False)
        
        # Format header
        ws = writer.sheets['All Variants']
        for cell in ws[1]:
            cell.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
            cell.font = Font(color="FFFFFF", bold=True)
        
        # Auto-width
        for i, col in enumerate(df.columns):
            width = min(max(df[col].astype(str).map(len).max(), len(col)) + 2, 40)
            ws.column_dimensions[chr(65 + i) if i < 26 else chr(64 + i//26) + chr(65 + i%26)].width = width
    
    logger.info(f"Exported {len(df)} rows to {path}")
    
    # Export errors
    if errors:
        err_path = os.path.join(OUTPUT_DIR, f"errors_{timestamp}.xlsx")
        pd.DataFrame(errors).to_excel(err_path, index=False)
    
    return path


# =============================================================================
# MAIN SCRAPER
# =============================================================================

def main():
    print("""
============================================================
     SPINNY ALL BRANDS SCRAPER - COMPLETE DATA EXTRACTION
============================================================
     Fetching: ALL 31 usable brands x Years 2000-2025
     Rate Limiting: 1.2s per request
     Checkpoint: Auto-save every 5 iterations
============================================================
    """)
    
    start_time = datetime.now()
    client = SpinnyClient()
    
    all_rows = []
    errors = []
    start_brand_idx = 0
    start_year_idx = 0
    
    # Check for resume
    checkpoint = load_checkpoint()
    if checkpoint:
        all_rows = checkpoint.get("rows", [])
        start_brand_idx = checkpoint.get("brand_idx", 0)
        start_year_idx = checkpoint.get("year_idx", 0)
        errors = checkpoint.get("errors", [])
        logger.info(f"Resuming: {len(all_rows)} rows, brand_idx={start_brand_idx}, year_idx={start_year_idx}")
    
    try:
        # Get ALL brands
        logger.info("Fetching all brands...")
        all_brands = client.get_all_brands()
        
        # Filter to usable brands only
        brands = [b for b in all_brands if b.get("is_usable", False)]
        logger.info(f"Found {len(brands)} usable brands out of {len(all_brands)} total")
        
        # Print all brand IDs for reference
        print("\n=== ALL USABLE BRANDS ===")
        for b in brands:
            print(f"  ID: {b['id']:3d} - {b['display_name']}")
        print(f"\nTotal: {len(brands)} brands\n")
        
        years = list(range(YEAR_START, YEAR_END + 1))
        total = len(brands) * len(years)
        
        with tqdm(total=total, desc="Scraping ALL Brands", unit="brand-year") as pbar:
            for brand_idx, brand in enumerate(brands):
                brand_id = brand["id"]
                brand_name = brand["display_name"]
                
                # Skip if resuming
                if brand_idx < start_brand_idx:
                    pbar.update(len(years))
                    continue
                
                for year_idx, year in enumerate(years):
                    # Skip if resuming within same brand
                    if brand_idx == start_brand_idx and year_idx < start_year_idx:
                        pbar.update(1)
                        continue
                    
                    pbar.set_description(f"{brand_name} ({year})")
                    
                    try:
                        models = client.get_models(brand_id, year)
                        
                        for model in models:
                            if not model.get("is_usable", True):
                                continue
                            
                            model_id = model.get("id")
                            model_name = model.get("display_name", model.get("name", ""))
                            body_type = model.get("body_type", "")
                            segment = model.get("segment_type", "")
                            
                            if not model_id:
                                continue
                            
                            try:
                                variants = client.get_variants(model_id, year)
                                
                                for v in variants:
                                    if not v.get("is_usable", True):
                                        continue
                                    
                                    row = normalize(v, brand_name, model_name, body_type, segment, year)
                                    all_rows.append(row)
                                    
                            except Exception as e:
                                errors.append({"brand": brand_name, "model": model_name, "year": year, "error": str(e)})
                    
                    except Exception as e:
                        errors.append({"brand": brand_name, "year": year, "error": str(e)})
                    
                    # Save checkpoint every 5 iterations
                    if pbar.n % 5 == 0 and all_rows:
                        save_checkpoint(all_rows, brand_idx, year_idx, errors)
                    
                    pbar.update(1)
        
        # Export to Excel
        output_path = export_excel(all_rows, errors)
        
        # Cleanup checkpoint
        delete_checkpoint()
        
        # Summary
        if all_rows:
            df = pd.DataFrame(all_rows)
            print(f"""
============================================================
                    SCRAPE COMPLETE!
============================================================
Total Variants:     {len(df):,}
Unique Brands:      {df['brand_name'].nunique()}
Unique Models:      {df['model_name'].nunique()}
Year Range:         {df['make_year'].min()} - {df['make_year'].max()}
Errors:             {len(errors)}
Output File:        {output_path}

Brands Scraped:
{chr(10).join(f"  - {b}" for b in sorted(df['brand_name'].unique()))}
============================================================
            """)
    
    except KeyboardInterrupt:
        save_checkpoint(all_rows, brand_idx, year_idx, errors)
        print("\n[INTERRUPTED] Progress saved. Run again to resume.")
    
    finally:
        client.close()
        print(f"\nTotal time: {datetime.now() - start_time}")


if __name__ == "__main__":
    main()
